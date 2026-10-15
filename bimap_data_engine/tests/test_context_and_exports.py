import io
import uuid

import frappe
from frappe.tests.utils import FrappeTestCase
from frappe.utils import cstr
from openpyxl import load_workbook

from bimap_data_engine.api.changes import preview_changes
from bimap_data_engine.api.metadata import get_schema
from bimap_data_engine.api.records import list_records
from bimap_data_engine.services.context import (
	create_context,
	validate_context_import_plan,
	validate_context_import_specification,
)
from bimap_data_engine.services.exports import build_job_results_csv, build_records_xlsx
from bimap_data_engine.services.jobs import create_job


class TestContextAndExports(FrappeTestCase):
	def setUp(self):
		frappe.set_user("Administrator")
		item_group = frappe.get_all("Item Group", filters={"is_group": 0}, pluck="name", limit=1)[0]
		stock_uom = frappe.get_all("UOM", pluck="name", limit=1)[0]
		self.item_code = f"BDE-CONTEXT-{uuid.uuid4().hex[:10].upper()}"
		self.item = frappe.get_doc(
			{
				"doctype": "Item",
				"item_code": self.item_code,
				"item_name": "BIMAP context item",
				"item_group": item_group,
				"stock_uom": stock_uom,
				"description": "Context original",
				"is_stock_item": 0,
			}
		).insert()
		self.context = create_context(
			"Item",
			"Edit",
			["description"],
			{"name": self.item_code},
			"bimap_document_engine",
		)

	def tearDown(self):
		frappe.set_user("Administrator")

	def _change(self, fieldname="description", value="Context update"):
		item = frappe.get_doc("Item", self.item_code)
		return {
			"operation": "set",
			"doctype": "Item",
			"name": item.name,
			"fieldname": fieldname,
			"original": item.get(fieldname),
			"value": value,
			"snapshot_modified": cstr(item.modified),
		}

	def test_context_narrows_schema_query_and_changes(self):
		context_id = self.context["context_id"]
		schema = get_schema("Item", context_id)
		self.assertEqual([field["fieldname"] for field in schema["fields"]], ["description"])
		rows = list_records(
			"Item", ["description"], {"name": self.item_code}, page_length=10, context_id=context_id
		)
		self.assertEqual([row.name for row in rows["rows"]], [self.item_code])
		self.assertTrue(preview_changes([self._change()], context_id)["valid"])
		with self.assertRaises(frappe.PermissionError):
			preview_changes([self._change("item_name", "Forbidden")], context_id)
		with self.assertRaises(frappe.PermissionError):
			list_records("Item", ["description"], {"name": ["!=", self.item_code]}, context_id=context_id)

	def test_context_cannot_widen_native_fields(self):
		with self.assertRaises(frappe.ValidationError):
			create_context("Item", "Edit", ["modified"], {}, "test")

	def test_import_context_restricts_specification_and_materialized_plan(self):
		context = create_context("Item", "Import", ["item_code"], {}, "test")
		context_id = context["context_id"]
		self.assertIn("/app/bimap-data-import?", context["route"])
		allowed_spec = {
			"nodes": [{
				"id": "items", "doctype": "Item", "sheet": "Items", "external_key": "Key",
				"mapping": [{"source": "Code", "target": "item_code"}],
			}]
		}
		validate_context_import_specification(context_id, allowed_spec)
		blocked_spec = frappe._dict(allowed_spec.copy())
		blocked_spec["nodes"] = [{**allowed_spec["nodes"][0], "defaults": {"description": "No"}}]
		with self.assertRaises(frappe.PermissionError):
			validate_context_import_specification(context_id, blocked_spec)
		allowed_plan = {
			"records": [{
				"node_id": "items", "external_key": "one", "doctype": "Item",
				"values": {"item_code": "BDE-CONTEXT-IMPORT"}, "children": {}, "references": [],
			}]
		}
		validate_context_import_plan(context_id, allowed_plan)
		allowed_plan["records"][0]["values"]["description"] = "No"
		with self.assertRaises(frappe.PermissionError):
			validate_context_import_plan(context_id, allowed_plan)

	def test_xlsx_export_honours_context_and_filters(self):
		content, count = build_records_xlsx(
			"Item",
			["description"],
			{"name": self.item_code},
			100,
			self.context["context_id"],
		)
		workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
		rows = list(workbook.active.iter_rows(values_only=True))
		workbook.close()
		self.assertEqual(count, 1)
		self.assertEqual(rows[0], ("Description",))
		self.assertEqual(rows[1], ("Context original",))

	def test_job_results_csv_contains_per_record_status(self):
		job_name = create_job([self._change()], batch_size=1)
		content = build_job_results_csv(job_name).decode("utf-8-sig")
		self.assertIn("sequence,status,doctype", content)
		self.assertIn("Pending,Item", content)

	def test_context_is_bound_to_owner(self):
		email = f"bde-context-{uuid.uuid4().hex[:10]}@example.invalid"
		frappe.get_doc(
			{
				"doctype": "User",
				"email": email,
				"first_name": "BDE context outsider",
				"send_welcome_email": 0,
				"user_type": "System User",
			}
		).insert(ignore_permissions=True)
		frappe.set_user(email)
		with self.assertRaises(frappe.PermissionError):
			get_schema("Item", self.context["context_id"])
