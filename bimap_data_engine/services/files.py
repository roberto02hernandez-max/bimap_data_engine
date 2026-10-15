"""Bounded CSV/XLSX parsing independent from the spreadsheet provider."""

import csv
import io
from pathlib import Path

import frappe
from openpyxl import load_workbook


MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_ROWS = 100_000
MAX_COLUMNS = 500


def _clean_cell(value):
	if value is None:
		return None
	if isinstance(value, str):
		return value.strip()
	return value


def _sheet_payload(title, rows):
	rows = list(rows)
	if not rows:
		return {"name": title, "headers": [], "rows": [], "row_count": 0}
	if len(rows) > MAX_ROWS + 1:
		frappe.throw(f"La hoja {title} excede {MAX_ROWS} filas.")
	width = max((len(row) for row in rows), default=0)
	if width > MAX_COLUMNS:
		frappe.throw(f"La hoja {title} excede {MAX_COLUMNS} columnas.")
	headers = []
	seen = set()
	for index, value in enumerate(rows[0], start=1):
		header = str(_clean_cell(value) or f"Column {index}")
		if header in seen:
			frappe.throw(f"La hoja {title} contiene el encabezado duplicado {header}.")
		seen.add(header)
		headers.append(header)
	data = []
	for raw in rows[1:]:
		values = [_clean_cell(value) for value in raw[: len(headers)]]
		values.extend([None] * (len(headers) - len(values)))
		if any(value not in (None, "") for value in values):
			data.append(dict(zip(headers, values, strict=True)))
	return {"name": title, "headers": headers, "rows": data, "row_count": len(data)}


def parse_file_bytes(filename, content):
	if isinstance(content, str):
		content = content.encode("utf-8")
	if not isinstance(content, bytes):
		frappe.throw("El contenido del archivo debe ser binario.")
	if len(content) > MAX_FILE_BYTES:
		frappe.throw(f"El archivo excede el límite de {MAX_FILE_BYTES // (1024 * 1024)} MB.")
	extension = Path(filename or "").suffix.lower()
	if extension == ".csv":
		text = content.decode("utf-8-sig")
		rows = list(csv.reader(io.StringIO(text)))
		return {"format": "CSV", "sheets": [_sheet_payload(Path(filename).stem or "CSV", rows)]}
	if extension in {".xlsx", ".xlsm"}:
		workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
		try:
			sheets = [_sheet_payload(sheet.title, sheet.iter_rows(values_only=True)) for sheet in workbook.worksheets]
		finally:
			workbook.close()
		return {"format": "XLSX", "sheets": sheets}
	frappe.throw("Formato no soportado. Usa CSV o XLSX.")


def parse_attached_file(file_url):
	file_name = frappe.db.get_value("File", {"file_url": file_url}, "name")
	if not file_name:
		frappe.throw("El archivo adjunto no existe.")
	file_doc = frappe.get_doc("File", file_name)
	file_doc.check_permission("read")
	content = file_doc.get_content()
	return parse_file_bytes(file_doc.file_name, content)
